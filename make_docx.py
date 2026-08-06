# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

SAFFRON="FF9933"; GREEN="138808"; GREEN_D="0E6606"
LIGHT_SAFF="FFF6EC"; LIGHT_GREEN="EAF6EA"; LINE="E7E7E7"
HINDI_FONT="Nirmala UI"
REPO="/projects/sandbox/Virali_insta_facebook_account"

doc=Document()
style=doc.styles["Normal"]; style.font.name=HINDI_FONT; style.font.size=Pt(10.5)
rpr=style.element.get_or_add_rPr()
rf=rpr.find(qn('w:rFonts'))
if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a),HINDI_FONT)

for s in doc.sections:
    s.top_margin=Mm(12); s.bottom_margin=Mm(12); s.left_margin=Mm(15); s.right_margin=Mm(15)

def cs(run,name=HINDI_FONT):
    rpr=run._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a),name)

def shade(p,fill):
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill)
    pPr.append(shd)

def lborder(p,color,size=32):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); l=OxmlElement('w:left')
    l.set(qn('w:val'),'single'); l.set(qn('w:sz'),str(size)); l.set(qn('w:space'),'6'); l.set(qn('w:color'),color)
    b.append(l); pPr.append(b)

def cell_bottom(cell,color=LINE,size=4):
    tcPr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders'); bot=OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),str(size)); bot.set(qn('w:space'),'0'); bot.set(qn('w:color'),color)
    borders.append(bot); tcPr.append(borders)

def sec(text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(5)
    shade(p,SAFFRON); lborder(p,GREEN,36)
    r=p.add_run(text); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); cs(r)

def bullets(items):
    for it in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3)
        if isinstance(it,tuple):
            r=p.add_run(it[0]); r.bold=True; cs(r); r2=p.add_run(it[1]); cs(r2)
        else:
            r=p.add_run(it); cs(r)

# ---------- HEADER (logo + party + title) ----------
logo=os.path.join(REPO,"bjp-logo.png")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(logo):
    p.add_run().add_picture(logo, width=Mm(20))
p.paragraph_format.space_after=Pt(0)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("भारतीय जनता पार्टी"); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor(0xe9,0x7e,0x12); cs(r)
p.paragraph_format.space_after=Pt(0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("जिला दतिया (मध्य प्रदेश)"); r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x6b,0x6b,0x6b); cs(r)
p.paragraph_format.space_after=Pt(2)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("परिचय — पत्र"); r.bold=True; r.font.size=Pt(22); r.underline=True; cs(r)
p.paragraph_format.space_after=Pt(6)

# ---------- OBJECTIVE ----------
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; shade(p,LIGHT_SAFF)
r=p.add_run("उद्देश्य: "); r.bold=True; cs(r)
r=p.add_run("एक समर्पित, अनुशासित एवं संघर्षशील कार्यकर्ता — जिसका एकमात्र लक्ष्य दतिया जिले में "
            "संगठन को बूथ-स्तर तक सशक्त कर पार्टी को पुनः विजयी बनाना है।"); r.bold=True; cs(r)
p.paragraph_format.space_after=Pt(8)

# ---------- INFO TABLE ----------
info=[("नाम","एडवोकेट कप्तान सिंह कुशवाह"),
      ("पिता का नाम","इंजीनियर रामस्वरूप सिंह कुशवाह (शासकीय सेवक, लोक निर्माण विभाग)"),
      ("जन्म तिथि","05 जनवरी 1993"),
      ("शैक्षणिक योग्यता","B.Sc (गणित), M.S.W (समाज कार्य), LL.B"),
      ("व्यवसाय","1. श्री गिरिराज जी प्राइवेट ITI कॉलेज, दतिया\n2. जिला एवं सत्र न्यायालय (डिस्ट्रिक्ट कोर्ट), दतिया — अधिवक्ता, 4 वर्ष"),
      ("निवास स्थान","रिछरा फाटक के बाहर, चिरईटोर माता मंदिर के पास, दतिया"),
      ("कार्यालय पता","सेवड़ा चुंगी मार्ग, चिरईटोर माता, दतिया"),
      ("मोबाइल","9584523005"),("ईमेल","singhdatkaptan@gmail.com")]
tbl=doc.add_table(rows=0,cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
tbl.columns[0].width=Mm(48); tbl.columns[1].width=Mm(132)
for label,val in info:
    cells=tbl.add_row().cells
    cells[0].width=Mm(48); cells[1].width=Mm(132)
    pl=cells[0].paragraphs[0]; rl=pl.add_run(label); rl.bold=True; rl.font.color.rgb=RGBColor(0x0e,0x66,0x06); cs(rl)
    pv=cells[1].paragraphs[0]
    for i,line in enumerate(val.split("\n")):
        if i>0: pv.add_run().add_break()
        rv=pv.add_run(line); rv.bold=True; cs(rv)
    cell_bottom(cells[0]); cell_bottom(cells[1])
doc.add_paragraph().paragraph_format.space_after=Pt(2)

# ---------- SPECIAL ----------
p=doc.add_paragraph(); shade(p,LIGHT_GREEN); lborder(p,GREEN,36)
r=p.add_run("🏆 विशेष उपलब्धि — बूथ क्रमांक 156"); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(0x0e,0x66,0x06); cs(r)
p2=doc.add_paragraph(); shade(p2,LIGHT_GREEN)
r2=p2.add_run("हाल के दतिया उपचुनाव में पूरे जिले में पराजय के बावजूद, मेरे प्रभार वाले बूथ क्रमांक 156 पर "
              "पार्टी को विजय/बढ़त प्राप्त हुई — जो मेरी बूथ-स्तरीय पकड़ का प्रमाण है।"); cs(r2)
p2.paragraph_format.space_after=Pt(6)

# ---------- SECTIONS ----------
sec("संगठनात्मक अनुभव (12 वर्ष का समर्पण)")
bullets([("वर्तमान दायित्व : ","जिला उपाध्यक्ष, भारतीय जनता युवा मोर्चा, दतिया"),
         ("पूर्व दायित्व : ","जिला मंत्री, भारतीय जनता युवा मोर्चा, दतिया (2020)"),
         ("चुनावी दायित्व : ","पोलिंग प्रभारी — बूथ क्रमांक 156"),
         "विगत 12 वर्षों से पार्टी में निरंतर सक्रिय एवं समर्पित कार्यकर्ता।",
         "मुख्यमंत्री, मंत्रीगण व विधायकों के दतिया आगमन पर निरंतर सहभागिता व समन्वय।",
         "बूथ स्तर से जिला स्तर तक संगठन को मज़बूत करने का व्यावहारिक अनुभव।"])

sec("संगठन के प्रति अटूट निष्ठा एवं समर्पण")
bullets(["पार्टी के हर दौर में — सुख हो या चुनौती — संगठन के प्रति पूर्ण निष्ठा के साथ निरंतर सक्रिय रहा।",
         "23 मई 2026 को माननीय डॉ. नरोत्तम मिश्रा जी द्वारा आयोजित कार्यक्रम में सक्रिय सहभागिता कर, संगठन को और अधिक सशक्त करने हेतु कार्य किया।",
         "23 मई 2026 को संगठन को मज़बूत करने का दृढ़ संकल्प लेकर बूथ-स्तर तक जनसंपर्क तेज़ किया।",
         "उपचुनाव में पार्टी के प्रत्याशी चयन का पूर्ण सम्मान करते हुए — व्यक्ति नहीं, बल्कि \"कमल\" एवं विचारधारा को सर्वोपरि रखते हुए — कठिन परिस्थितियों में भी पूरी शक्ति से प्रचार किया।",
         "पार्टी हित सर्वोपरि — यही मेरी कार्यशैली एवं पहचान है।"])

sec("वैचारिक एवं सामाजिक पृष्ठभूमि")
bullets(["राष्ट्रीय स्वयंसेवक संघ (RSS) के प्रत्येक कार्यक्रम/आयोजन में निरंतर सहभागिता।",
         "M.S.W एवं अधिवक्ता होने के नाते जनसेवा व निःशुल्क कानूनी सहायता में सक्रिय।",
         "राष्ट्रवादी विचारधारा, पार्टी अनुशासन व संगठन के प्रति पूर्ण निष्ठा।"])

sec("दतिया जिले हेतु संकल्प / विज़न")
bullets(["बूथ-स्तर तक संगठन को पुनः सशक्त कर हर कार्यकर्ता को सक्रिय व सम्मानित करना।",
         "पार्टी की नीतियों व जनकल्याणकारी योजनाओं को घर-घर पहुँचाना।",
         "युवाओं व समाज के हर वर्ग को जोड़कर दतिया को पार्टी का मज़बूत गढ़ बनाना।"])

sec("सोशल मीडिया / जनसंपर्क (अनेक पोस्ट व रील वायरल)")
bullets(["Instagram : @adv.KaptanSingh","Facebook : adv.KaptanSingh1","Facebook Page : adv.KaptanSingh",
         "Twitter (X) : @adv.KaptanSingh","Threads : @adv.KaptanSingh"])
p=doc.add_paragraph(); r=p.add_run("रुचि : "); r.bold=True; r.font.color.rgb=RGBColor(0x0e,0x66,0x06); cs(r)
r2=p.add_run("समाज सेवा, जनसंपर्क एवं संगठनात्मक कार्य"); cs(r2)

# ---------- PAGE BREAK -> PHOTO APPENDIX ----------
doc.add_page_break()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("फोटो परिशिष्ट"); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(0xe9,0x7e,0x12); cs(r)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("वरिष्ठ नेतृत्व, संगठन एवं सामाजिक कार्यों के साथ सहभागिता"); r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x6b,0x6b,0x6b); cs(r)
p.paragraph_format.space_after=Pt(6)

caps=["माननीय मुख्यमंत्री जी के साथ","माननीय डॉ. नरोत्तम मिश्रा जी के साथ","विधायक / मंत्री जी के साथ",
      "उपचुनाव प्रचार के दौरान","संगठन बैठक / रैली","सामाजिक सेवा कार्य"]
pt=doc.add_table(rows=2,cols=3); pt.style="Table Grid"; pt.alignment=WD_TABLE_ALIGNMENT.CENTER
idx=0
for row in pt.rows:
    row.height=Mm(34)
    for c in row.cells:
        pp=c.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=pp.add_run(f"[ फोटो {idx+1} ]"); rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(0x9a,0x9a,0x9a); cs(rr)
        cp=c.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cr=cp.add_run(caps[idx]); cr.font.size=Pt(9); cr.bold=True; cr.font.color.rgb=RGBColor(0x55,0x55,0x55); cs(cr)
        idx+=1

doc.add_paragraph().paragraph_format.space_after=Pt(2)
sec("वायरल पोस्ट / रील")
vt=doc.add_table(rows=1,cols=6); vt.style="Table Grid"; vt.alignment=WD_TABLE_ALIGNMENT.CENTER
vl=["पोस्ट/रील 1","पोस्ट/रील 2","पोस्ट/रील 3","पोस्ट/रील 4","पोस्ट/रील 5","QR कोड"]
vt.rows[0].height=Mm(36)
for i,c in enumerate(vt.rows[0].cells):
    pp=c.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=pp.add_run(vl[i]); rr.font.size=Pt(9); rr.bold=(i==5)
    rr.font.color.rgb=RGBColor(0x0e,0x66,0x06) if i==5 else RGBColor(0x9a,0x9a,0x9a); cs(rr)

# ---------- SIGN OFF ----------
doc.add_paragraph(); doc.add_paragraph()
st=doc.add_table(rows=1,cols=2); st.alignment=WD_TABLE_ALIGNMENT.CENTER
lc=st.rows[0].cells[0].paragraphs[0]; r=lc.add_run("दिनांक: ____________\nस्थान: दतिया"); r.bold=True; cs(r)
rc=st.rows[0].cells[1].paragraphs[0]; rc.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=rc.add_run("________________________\nएडवोकेट कप्तान सिंह कुशवाह\nमोबाइल: 9584523005"); r.bold=True; cs(r)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(14)
r=p.add_run("जय हिन्द • जय भाजपा"); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor(0xe9,0x7e,0x12); cs(r)

out=os.path.join(REPO,"Bio-Data-Kaptan-Singh.docx"); doc.save(out); print("Saved styled DOCX ->",out)
